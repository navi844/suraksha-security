from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
from typing import Any

import fitz
from bs4 import BeautifulSoup
from docx import Document as DocxDocument

from pii import redact_text
from schemas import SourceDocument


REQUIRED_PII_HINTS = {'email', 'phone'}


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode('utf-8')).hexdigest()


def load_pdf(path: Path) -> list[SourceDocument]:
    docs: list[SourceDocument] = []
    with fitz.open(path) as document:
        if document.page_count == 0:
            raise ValueError(f'PDF document is empty: {path}')
        for page_index in range(document.page_count):
            page = document[page_index]
            text = page.get_text()
            if not text.strip():
                continue
            record = SourceDocument(
                text=text,
                source=str(path.name),
                document_id='SUR-ACP-001',
                version='1.0',
                effective_date='2026-01-15',
                file_format='pdf',
                locator_type='page',
                locator_value=str(page_index + 1),
                page=page_index + 1,
                section='content',
                content_hash=sha256_text(text),
            )
            docs.append(record)
    if not docs:
        raise ValueError(f'No extractable text found in PDF: {path}')
    return docs


def load_docx(path: Path) -> list[SourceDocument]:
    document = DocxDocument(str(path))
    sections: list[str] = []
    current = []
    for paragraph in document.paragraphs:
        text = (paragraph.text or '').strip()
        if not text:
            continue
        if text.startswith('Section '):
            if current:
                sections.append('\n'.join(current))
            current = [text]
        else:
            current.append(text)
    if current:
        sections.append('\n'.join(current))

    if not sections:
        raise ValueError(f'DOCX document is empty: {path}')

    records: list[SourceDocument] = []
    for idx, section in enumerate(sections):
        record = SourceDocument(
            text=section,
            source=str(path.name),
            document_id='SUR-IRP-002',
            version='1.2',
            effective_date='2026-03-01',
            file_format='docx',
            locator_type='section',
            locator_value=f'section_{idx + 1}',
            section=f'section_{idx + 1}',
            content_hash=sha256_text(section),
        )
        records.append(record)
    return records


def load_html(path: Path) -> list[SourceDocument]:
    text = path.read_text(encoding='utf-8')
    soup = BeautifulSoup(text, 'html.parser')
    for tag in soup.find_all(['script', 'style', 'nav', 'header', 'footer']):
        tag.decompose()
    main = soup.find('main') or soup
    sections: list[SourceDocument] = []
    question_blocks = main.find_all(['section', 'h2', 'h3', 'p', 'blockquote', 'pre'])
    content_parts: list[str] = []
    for element in question_blocks:
        text_value = element.get_text(' ', strip=True)
        if not text_value:
            continue
        content_parts.append(text_value)
    combined = '\n'.join(content_parts)
    if not combined.strip():
        raise ValueError(f'HTML document is empty: {path}')
    record = SourceDocument(
        text=combined,
        source=str(path.name),
        document_id='SUR-SAF-003',
        version='2.0',
        effective_date='2026-05-10',
        file_format='html',
        locator_type='html',
        locator_value='body',
        section='faq',
        content_hash=sha256_text(combined),
    )
    sections.append(record)
    return sections


def load_csv(path: Path) -> list[SourceDocument]:
    with path.open('r', encoding='utf-8', newline='') as handle:
        lines = handle.readlines()

    cleaned = []
    for line in lines:
        if line.strip().startswith('Document ID:') or line.strip().startswith('Version:') or line.strip().startswith('Effective Date:'):
            continue
        cleaned.append(line)

    reader = csv.DictReader(cleaned)
    required = ['incident_id', 'date', 'category', 'severity', 'status', 'summary']
    if reader.fieldnames is None or any(field not in (reader.fieldnames or []) for field in required):
        raise ValueError(f'Missing required CSV columns in {path}')

    rows: list[SourceDocument] = []
    for row_number, row in enumerate(reader, start=1):
        if not row or not any((value or '').strip() for value in row.values()):
            continue
        text = ' | '.join(f'{key}: {value}' for key, value in row.items() if value is not None)
        records = SourceDocument(
            text=text,
            source=str(path.name),
            document_id='SUR-SIL-004',
            version='1.0',
            effective_date='2026-09-30',
            file_format='csv',
            locator_type='row',
            locator_value=str(row_number),
            row_number=row_number,
            record_id=row.get('incident_id', f'row-{row_number}'),
            section='incident_record',
            content_hash=sha256_text(text),
        )
        rows.append(records)
    if not rows:
        raise ValueError(f'CSV document contains no valid rows: {path}')
    return rows


def load_json(path: Path) -> list[SourceDocument]:
    payload = json.loads(path.read_text(encoding='utf-8'))
    if 'vendor_categories' not in payload:
        raise ValueError(f'Missing vendor_categories in JSON document: {path}')
    categories = payload['vendor_categories']
    if not isinstance(categories, list) or not categories:
        raise ValueError(f'No vendor categories in JSON document: {path}')

    rows: list[SourceDocument] = []
    for index, category in enumerate(categories):
        if 'category_name' not in category:
            raise ValueError(f'Vendor category missing category_name in {path}')
        text = json.dumps(category, sort_keys=True)
        records = SourceDocument(
            text=text,
            source=str(path.name),
            document_id=payload.get('document_id', 'SUR-AVR-005'),
            version=payload.get('version', '1.0'),
            effective_date=payload.get('effective_date', '2026-06-01'),
            file_format='json',
            locator_type='category',
            locator_value=category['category_name'],
            section='vendor_category',
            record_id=category['category_name'],
            content_hash=sha256_text(text),
        )
        rows.append(records)
    return rows


def load_document(path: Path) -> list[SourceDocument]:
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        return load_pdf(path)
    if suffix == '.docx':
        return load_docx(path)
    if suffix == '.html':
        return load_html(path)
    if suffix == '.csv':
        return load_csv(path)
    if suffix == '.json':
        return load_json(path)
    raise ValueError(f'Unsupported document format: {path}')


def load_directory(path: Path) -> list[SourceDocument]:
    if not path.exists():
        raise FileNotFoundError(f'Documents directory does not exist: {path}')
    all_files = sorted(path.iterdir())
    documents: list[SourceDocument] = []
    seen_hashes: set[str] = set()
    for file_path in all_files:
        if file_path.is_file() and file_path.suffix.lower() in {'.pdf', '.docx', '.html', '.csv', '.json'}:
            loaded = load_document(file_path)
            for record in loaded:
                if record.content_hash in seen_hashes:
                    continue
                seen_hashes.add(record.content_hash)
                records = record.model_copy(update={'pii_redacted': False, 'pii_types': [], 'injection_indicator_present': False})
                documents.append(records)
    return documents
